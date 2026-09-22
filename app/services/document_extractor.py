import fitz  # PyMuPDF
import docx
from docx.text.paragraph import Paragraph
from docx.table import Table
import pandas as pd
from pathlib import Path
from typing import List
from pydantic import BaseModel


class ExtractedBlock(BaseModel):
    type: str  # "page", "paragraph", "table", "document", "row"
    index: int  # 1-based index
    text: str


class ExtractedDocument(BaseModel):
    document_id: str
    source_filename: str
    file_type: str
    blocks: List[ExtractedBlock]


def extract_pdf(file_path: Path) -> List[ExtractedBlock]:
    """Extracts text page-by-page from a PDF file."""
    blocks = []
    try:
        doc = fitz.open(file_path)
        for page_idx, page in enumerate(doc):
            page_num = page_idx + 1
            text = page.get_text()
            blocks.append(
                ExtractedBlock(type="page", index=page_num, text=text)
            )
        doc.close()
    except Exception as e:
        raise ValueError(f"Failed to extract PDF: {str(e)}")
    return blocks


def extract_docx(file_path: Path) -> List[ExtractedBlock]:
    """Extracts paragraphs and tables from a DOCX file preserving reading order."""
    blocks = []
    try:
        doc = docx.Document(file_path)
        paragraph_idx = 1
        table_idx = 1

        body_elements = doc.element.body
        for child in body_elements.iterchildren():
            tag = child.tag
            if tag.endswith("p"):
                p = Paragraph(child, doc)
                text = p.text.strip()
                if text:
                    blocks.append(
                        ExtractedBlock(
                            type="paragraph", index=paragraph_idx, text=text
                        )
                    )
                    paragraph_idx += 1
            elif tag.endswith("tbl"):
                tbl = Table(child, doc)
                table_lines = []
                for row_idx, row in enumerate(tbl.rows):
                    row_cells = [cell.text.strip() for cell in row.cells]
                    row_text = " | ".join(row_cells)
                    table_lines.append(f"Row {row_idx + 1}: {row_text}")

                table_text = "\n".join(table_lines)
                blocks.append(
                    ExtractedBlock(type="table", index=table_idx, text=table_text)
                )
                table_idx += 1
    except Exception as e:
        raise ValueError(f"Failed to extract DOCX: {str(e)}")
    return blocks


def extract_txt(file_path: Path) -> List[ExtractedBlock]:
    """Extracts text from a UTF-8 text file with fallback decoding."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()
        return [ExtractedBlock(type="document", index=1, text=text)]
    except Exception as e:
        raise ValueError(f"Failed to extract TXT: {str(e)}")


def extract_csv(file_path: Path) -> List[ExtractedBlock]:
    """Extracts rows from a CSV file converting each to structured text."""
    try:
        df = pd.read_csv(file_path, keep_default_na=False)
        if df.empty:
            raise ValueError("CSV file has no data rows.")

        blocks = []
        for row_idx, row in df.iterrows():
            row_items = []
            for col in df.columns:
                val = str(row[col]).strip()
                row_items.append(f"{col}: {val}")
            row_text = " | ".join(row_items)
            blocks.append(
                ExtractedBlock(type="row", index=row_idx + 1, text=row_text)
            )
        return blocks
    except pd.errors.EmptyDataError:
        raise ValueError("CSV file is empty.")
    except pd.errors.ParserError:
        raise ValueError("Malformed CSV formatting error.")
    except UnicodeDecodeError:
        raise ValueError("CSV encoding error: failed to decode file bytes.")
    except Exception as e:
        # Avoid exposing raw stack details
        raise ValueError(f"Failed to parse CSV document: {str(e)}")


def extract_document(
    document_id: str, original_filename: str, file_path: Path, file_type: str
) -> ExtractedDocument:
    """Dispatches document extraction based on file type."""
    if not file_path.exists():
        raise FileNotFoundError("Raw document file not found on disk.")

    file_type_lower = file_type.lower().strip()

    if file_type_lower == "pdf":
        blocks = extract_pdf(file_path)
    elif file_type_lower == "docx":
        blocks = extract_docx(file_path)
    elif file_type_lower == "txt":
        blocks = extract_txt(file_path)
    elif file_type_lower == "csv":
        blocks = extract_csv(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")

    return ExtractedDocument(
        document_id=document_id,
        source_filename=original_filename,
        file_type=file_type_lower,
        blocks=blocks,
    )
