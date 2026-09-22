import os
import tempfile
import json
from pathlib import Path
import pytest
import fitz
import docx
from fastapi.testclient import TestClient
from app.main import app
from app.config import get_settings, Settings
from app.services.document_extractor import extract_document

client = TestClient(app)


@pytest.fixture
def temp_upload_dir():
    with tempfile.TemporaryDirectory() as tmpdir:
        yield Path(tmpdir)


@pytest.fixture(autouse=True)
def override_settings(temp_upload_dir, monkeypatch):
    monkeypatch.setenv("UPLOAD_DIR", str(temp_upload_dir))
    monkeypatch.setenv("MAX_UPLOAD_SIZE", "10000000")

    get_settings.cache_clear()
    yield
    get_settings.cache_clear()


def test_pdf_extraction_integration(temp_upload_dir):
    # 1. Create a dummy multi-page PDF
    pdf_path = temp_upload_dir / "test_doc.pdf"
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((50, 50), "Hello on Page 1")
    page2 = doc.new_page()
    page2.insert_text((50, 50), "Greetings on Page 2")
    doc.save(str(pdf_path))
    doc.close()

    # 2. Upload it
    with open(pdf_path, "rb") as f:
        resp = client.post(
            "/documents/upload",
            files={"file": ("test_doc.pdf", f, "application/pdf")},
        )
    assert resp.status_code == 200
    doc_id = resp.json()["document_id"]

    # 3. Call extract endpoint
    extract_resp = client.post(f"/documents/{doc_id}/extract")
    assert extract_resp.status_code == 200
    data = extract_resp.json()

    assert data["document_id"] == doc_id
    assert data["source_filename"] == "test_doc.pdf"
    assert data["file_type"] == "pdf"
    assert len(data["blocks"]) == 2

    assert data["blocks"][0]["type"] == "page"
    assert data["blocks"][0]["index"] == 1
    assert "Page 1" in data["blocks"][0]["text"]

    assert data["blocks"][1]["type"] == "page"
    assert data["blocks"][1]["index"] == 2
    assert "Page 2" in data["blocks"][1]["text"]


def test_docx_extraction_reading_order_integration(temp_upload_dir):
    # 1. Create a docx with paragraph -> table -> paragraph
    docx_path = temp_upload_dir / "test_doc.docx"
    doc = docx.Document()
    doc.add_paragraph("Intro Paragraph")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    table.cell(1, 0).text = "A2"
    table.cell(1, 1).text = "B2"

    doc.add_paragraph("Outro Paragraph")
    doc.save(str(docx_path))

    # 2. Upload it
    with open(docx_path, "rb") as f:
        resp = client.post(
            "/documents/upload",
            files={
                "file": (
                    "test_doc.docx",
                    f,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
    assert resp.status_code == 200
    doc_id = resp.json()["document_id"]

    # 3. Extract and check reading order
    extract_resp = client.post(f"/documents/{doc_id}/extract")
    assert extract_resp.status_code == 200
    data = extract_resp.json()

    blocks = data["blocks"]
    assert len(blocks) == 3

    # Paragraph 1
    assert blocks[0]["type"] == "paragraph"
    assert blocks[0]["index"] == 1
    assert blocks[0]["text"] == "Intro Paragraph"

    # Table 1
    assert blocks[1]["type"] == "table"
    assert blocks[1]["index"] == 1
    assert "Row 1: A1 | B1" in blocks[1]["text"]
    assert "Row 2: A2 | B2" in blocks[1]["text"]

    # Paragraph 2 (index should be 2)
    assert blocks[2]["type"] == "paragraph"
    assert blocks[2]["index"] == 2
    assert blocks[2]["text"] == "Outro Paragraph"


def test_txt_extraction_integration(temp_upload_dir):
    # 1. Create TXT file with utf-8 encoding and invalid characters to test fallback
    txt_path = temp_upload_dir / "test.txt"
    # Write some bytes containing invalid utf-8 sequences along with valid text
    with open(txt_path, "wb") as f:
        f.write(b"Hello world. \xff Invalid character test.")

    # 2. Upload
    with open(txt_path, "rb") as f:
        resp = client.post(
            "/documents/upload", files={"file": ("test.txt", f, "text/plain")}
        )
    assert resp.status_code == 200
    doc_id = resp.json()["document_id"]

    # 3. Extract
    extract_resp = client.post(f"/documents/{doc_id}/extract")
    assert extract_resp.status_code == 200
    data = extract_resp.json()

    assert len(data["blocks"]) == 1
    assert data["blocks"][0]["type"] == "document"
    assert "Hello world" in data["blocks"][0]["text"]
    # Verify fallback unicode replacement character is present due to errors='replace'
    assert "\ufffd" in data["blocks"][0]["text"]


def test_csv_extraction_integration(temp_upload_dir):
    # 1. Create a normal CSV
    csv_path = temp_upload_dir / "test.csv"
    with open(csv_path, "w", encoding="utf-8") as f:
        f.write("Name,Age,Role\nVivek,25,Engineer\nShinde,30,Manager")

    # 2. Upload
    with open(csv_path, "rb") as f:
        resp = client.post(
            "/documents/upload", files={"file": ("test.csv", f, "text/csv")}
        )
    assert resp.status_code == 200
    doc_id = resp.json()["document_id"]

    # 3. Extract
    extract_resp = client.post(f"/documents/{doc_id}/extract")
    assert extract_resp.status_code == 200
    data = extract_resp.json()

    blocks = data["blocks"]
    assert len(blocks) == 2
    assert blocks[0]["type"] == "row"
    assert blocks[0]["index"] == 1
    assert "Name: Vivek | Age: 25 | Role: Engineer" in blocks[0]["text"]

    assert blocks[1]["type"] == "row"
    assert blocks[1]["index"] == 2
    assert "Name: Shinde | Age: 30 | Role: Manager" in blocks[1]["text"]


def test_csv_extraction_error_scenarios(temp_upload_dir):
    # Test Empty CSV / Malformed CSV by manually saving metadata JSON to trigger extract errors
    doc_id = "00000000-0000-0000-0000-000000000001"

    # Create an empty CSV file
    empty_csv_path = temp_upload_dir / "empty.csv"
    empty_csv_path.touch()

    meta = {
        "document_id": doc_id,
        "original_filename": "empty.csv",
        "stored_filename": "empty.csv",
        "file_type": "csv",
        "file_size": 0,
        "status": "uploaded",
    }
    with open(temp_upload_dir / f"{doc_id}.json", "w") as f:
        json.dump(meta, f)

    # Call endpoint - should return 422 because CSV parsing fails or is empty
    resp = client.post(f"/documents/{doc_id}/extract")
    assert resp.status_code == 422
    assert "empty" in resp.json()["detail"].lower()


def test_endpoint_validation_uuid_format():
    # 1. Invalid UUID format should return 400 Bad Request
    resp = client.post("/documents/invalid-uuid-format/extract")
    assert resp.status_code == 400
    assert "Invalid document ID format" in resp.json()["detail"]


def test_endpoint_missing_document():
    # 2. Document not found on metadata lookup should return 404
    missing_doc_id = "11111111-1111-1111-1111-111111111111"
    resp = client.post(f"/documents/{missing_doc_id}/extract")
    assert resp.status_code == 404
    assert "metadata not found" in resp.json()["detail"].lower()
